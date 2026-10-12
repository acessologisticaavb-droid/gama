"""Verify flat tabular Excel export of Dados Gerais."""
import io, os, random, string
import requests
from openpyxl import Workbook, load_workbook
BASE = os.popen("grep REACT_APP_BACKEND_URL /app/frontend/.env | cut -d= -f2").read().strip() + "/api"
RID = "".join(random.choices(string.ascii_uppercase + string.digits, k=5))
def login(s,u,p): r=s.post(f"{BASE}/auth/login",json={"username":u,"password":p}); r.raise_for_status(); return r.json()
sa=requests.Session(); login(sa,"admin","admin123")
c=sa.post(f"{BASE}/companies",json={"nome":f"Export {RID}","nome_fantasia":"E","cnpj":f"{random.randint(10,99)}.{random.randint(100,999)}.{random.randint(100,999)}/0001-{random.randint(10,99)}","ramo":"X","logo":None}).json()
cid=c["id"]; key=sa.get(f"{BASE}/companies/{cid}/key").json()["key"]
al=f"adm_{RID.lower()}"; sa.post(f"{BASE}/companies/{cid}/admin",json={"nome":"A","funcao":"G","login":al,"senha":"Senha@123"})
ca=requests.Session(); login(ca,al,"Senha@123"); ca.post(f"{BASE}/auth/validate-key",json={"key":key})
ca.post(f"{BASE}/inv-groups",json={"nome":"Inventário Geral"})
iid=ca.post(f"{BASE}/inventories",json={"tipo":"Inventário Geral","grupo_produtos":"FIO MAQUINA","data":"2026-10-10","periodo":"10/2026"}).json()["id"]
wb=Workbook(); ws=wb.active
ws.append(["Material","Texto breve material","Lote","UM básica","Pos.depósito","TipoEstoq.","Unidade de depósito","Dias em estoque","Estoque disponível","Peso Bruto"])
for s in ["1001519427","1001516821","1001519745"]:
    ws.append([f"80-00{s[-3:]}",f"Vergalhao {s[-3:]}",f"K{s[-4:]}","UN","GIB-B03","01", s, 14, 500, 2.021])
buf=io.BytesIO(); wb.save(buf); buf.seek(0)
ca.post(f"{BASE}/inventories/{iid}/import-cutoff?commit=true",files={"file":("c.xlsx",buf.read(),"application/octet-stream")})
wb2=Workbook(); ws2=wb2.active; ws2.append(["Cód. de Barras","Posição Física"])
for s in ["1001519427","1001516821"]: ws2.append([f"{s}000X00R","GIB-B03"])
buf2=io.BytesIO(); wb2.save(buf2); buf2.seek(0)
ca.post(f"{BASE}/inventories/{iid}/import-physical?commit=true",files={"file":("p.xlsx",buf2.read(),"application/octet-stream")})
xl=ca.get(f"{BASE}/inventories/{iid}/export-excel")
wbx=load_workbook(io.BytesIO(xl.content)); wsx=wbx.active
rows=list(wsx.iter_rows(values_only=True))
print("HEADER:", rows[0][:9])
print("ROW2  :", rows[1][:9])
print("TOTAL DATA ROWS:", len(rows)-1)
# header must be first row (no institutional block)
assert rows[0][0]=="Empresa" and rows[0][2]=="Inventário"
assert rows[1][0]=="Export "+RID  # empresa repeated per row
assert rows[1][2].startswith("INV-")
assert len(rows)-1 == 3  # 2 ok + 1 faltante
sa.delete(f"{BASE}/companies/{cid}?confirm={requests.utils.quote(c['nome'])}")
print("\n✅ EXCEL TABULAR OK (linha1=cabeçalhos, metadados repetidos por linha, faltante incluído)")
