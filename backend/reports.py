"""Excel export and PDF report generation."""
import io
import base64
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.piecharts import Pie

INV_META_COLUMNS = [
    ("empresa", "Empresa"),
    ("cnpj", "CNPJ"),
    ("inventario", "Inventário"),
    ("tipo", "Tipo de Inventário"),
    ("grupo_produtos", "Grupo de Produtos"),
    ("data", "Data do Inventário"),
    ("periodo", "Período"),
]

DADOS_GERAIS_COLUMNS = [
    ("sku_sistema", "SKU Sistema"),
    ("cod_material", "Código do Material"),
    ("descricao", "Descrição"),
    ("lote", "Lote"),
    ("um", "Unidade de Medida"),
    ("tipo_estoque", "Tipo de Estoque"),
    ("pos_sistema", "Posição Sistema"),
    ("dias_estoque", "Dias em Estoque"),
    ("peso", "Peso"),
    ("sku_fisico", "SKU Físico"),
    ("pos_fisica", "Posição Física"),
    ("cod_barras", "Código de Barras"),
    ("validacao_posicao", "Validação Posição"),
    ("status_inventario", "Status Inventario"),
    ("validacao_sku", "Validação SKU"),
    ("gatilho", "Gatilho"),
    ("motivo", "Motivo"),
    ("resolucao", "Resolução"),
    ("resolucao_status", "Status da Resolução"),
    ("data_resolucao", "Data de Conclusão"),
    ("concluido_por", "Usuário que concluiu"),
]

ALL_COLUMNS = INV_META_COLUMNS + DADOS_GERAIS_COLUMNS

STATUS_DISPLAY = {
    "CONTAGEM OK": "SKU no fisico", "SKU FALTANTE": "SKU faltante",
    "SKU FORA DO SISTEMA": "SKU fora do sistema", "UD DUPLICADA": "UD duplicada",
    "SEM ETIQUETA": "Sem etiqueta",
}


def build_excel(company, inventory, records) -> bytes:
    """Flat, normalized sheet: row 1 = headers, each record = 1 row, inv metadata repeated per row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados Gerais"
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F2937")

    meta_values = {
        "empresa": company.get("nome", ""),
        "cnpj": company.get("cnpj", ""),
        "inventario": inventory.get("codigo", ""),
        "tipo": inventory.get("tipo", ""),
        "grupo_produtos": inventory.get("grupo_produtos", ""),
        "data": inventory.get("data", ""),
        "periodo": inventory.get("periodo", ""),
    }

    # Row 1 = headers
    for ci, (_, label) in enumerate(ALL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=ci, value=label)
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    r = 2
    for rec in records:
        for ci, (key, _) in enumerate(ALL_COLUMNS, start=1):
            if key in meta_values:
                val = meta_values[key]
            elif key == "status_inventario":
                val = STATUS_DISPLAY.get(rec.get(key), rec.get(key, ""))
            else:
                val = rec.get(key, "")
            ws.cell(row=r, column=ci, value="" if val is None else str(val))
        r += 1

    for ci in range(1, len(ALL_COLUMNS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = 18
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bar_chart(data, labels, title):
    d = Drawing(240, 150)
    bc = VerticalBarChart()
    bc.x = 20
    bc.y = 20
    bc.height = 110
    bc.width = 200
    bc.data = [data]
    bc.categoryAxis.categoryNames = labels
    bc.valueAxis.valueMin = 0
    bc.bars[0].fillColor = colors.HexColor("#3B82F6")
    bc.categoryAxis.labels.fontSize = 6
    bc.categoryAxis.labels.angle = 20
    bc.categoryAxis.labels.dy = -6
    d.add(bc)
    return d


def _pie_chart(data, labels):
    d = Drawing(200, 150)
    pie = Pie()
    pie.x = 60
    pie.y = 20
    pie.width = 110
    pie.height = 110
    pie.data = data
    pie.labels = [str(x) for x in labels]
    palette = ["#10B981", "#F59E0B", "#3B82F6", "#EF4444", "#06B6D4", "#8B5CF6"]
    for i in range(len(data)):
        pie.slices[i].fillColor = colors.HexColor(palette[i % len(palette)])
    d.add(pie)
    return d


def build_pdf(company, inventory, kpis, resumo_material, pendencias, trigger_kpis) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=15 * mm, bottomMargin=15 * mm,
                            leftMargin=14 * mm, rightMargin=14 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#111827"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, textColor=colors.HexColor("#1F2937"))
    normal = styles["Normal"]
    elems = []

    # Header — logo preserved with aspect ratio (object-fit: contain), no artificial background
    header_cells = []
    logo = company.get("logo")
    if logo and isinstance(logo, str) and logo.startswith("data:image"):
        try:
            b64 = logo.split(",", 1)[1]
            raw = base64.b64decode(b64)
            ir = ImageReader(io.BytesIO(raw))
            iw, ih = ir.getSize()
            max_w, max_h = 38 * mm, 24 * mm
            ratio = min(max_w / iw, max_h / ih)
            img = Image(io.BytesIO(raw), width=iw * ratio, height=ih * ratio)
            img.hAlign = "LEFT"
            header_cells.append(img)
        except Exception:
            header_cells.append(Paragraph("GAMA", title_style))
    else:
        header_cells.append(Paragraph("GAMA", title_style))
    info = (
        f"<b>GAMA — Gestão Industrial</b><br/>"
        f"<b>{company.get('nome','')}</b><br/>CNPJ: {company.get('cnpj','')}<br/>"
        f"Inventário: <b>{inventory.get('codigo','')}</b>"
    )
    header_tbl = Table([[header_cells[0], Paragraph(info, normal)]], colWidths=[35 * mm, 140 * mm])
    header_tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    elems.append(header_tbl)
    elems.append(Spacer(1, 6))
    elems.append(Paragraph(
        f"Tipo: {inventory.get('tipo','')} &nbsp;|&nbsp; Grupo: {inventory.get('grupo_produtos','')} "
        f"&nbsp;|&nbsp; Data: {inventory.get('data','')} &nbsp;|&nbsp; Período: {inventory.get('periodo','')}",
        normal))
    elems.append(Spacer(1, 10))

    # KPIs
    elems.append(Paragraph("Indicadores (KPIs)", h2))
    kpi_data = [
        ["UDs no Sistema", kpis["cutoff_total"], "UDs Coletadas", kpis["physical_total"]],
        ["UDs Faltantes", kpis["faltantes"], "SKUs Fora do Sistema", kpis["fora_sistema"]],
        ["UDs Duplicadas", kpis["duplicadas"], "Sem Etiquetas", kpis["sem_etiqueta"]],
        ["% Posição Correta", f'{kpis["perc_posicao"]}%', "% Acuracidade", f'{kpis["perc_acuracidade"]}%'],
    ]
    kt = Table(kpi_data, colWidths=[45 * mm, 42 * mm, 45 * mm, 42 * mm])
    kt.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F4F6")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F3F4F6")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
    ]))
    elems.append(kt)
    elems.append(Spacer(1, 12))

    # Charts — relatório gerencial visual
    elems.append(Paragraph("Gráficos", h2))
    row1 = Table([[
        _bar_chart([kpis["cutoff_total"], kpis["physical_total"]], ["Sistema", "Coletadas"], "UDs"),
        _pie_chart([max(kpis["posicao_correta"], 0.0001), max(kpis["posicao_errada"], 0.0001)], ["Correta", "Errada"]),
    ]], colWidths=[90 * mm, 85 * mm])
    elems.append(row1)
    elems.append(Spacer(1, 6))
    pend_total = kpis["faltantes"] + kpis["fora_sistema"] + kpis["duplicadas"] + kpis["sem_etiqueta"]
    row2 = Table([[
        _pie_chart([max(kpis["contagem_ok"], 0.0001), max(pend_total, 0.0001)], ["OK", "Pend."]),
        _bar_chart(list(trigger_kpis.values()) or [0], list(trigger_kpis.keys()) or ["—"], "Gatilhos"),
    ]], colWidths=[85 * mm, 90 * mm])
    elems.append(row2)
    elems.append(Spacer(1, 12))

    # Trigger KPIs
    if trigger_kpis:
        elems.append(Paragraph("Pendências por Gatilho", h2))
        td = [["Gatilho", "Ocorrências"]] + [[k, str(v)] for k, v in trigger_kpis.items()]
        tt = Table(td, colWidths=[120 * mm, 40 * mm])
        tt.setStyle(_table_style())
        elems.append(tt)
        elems.append(Spacer(1, 12))

    # Resumo por material
    elems.append(Paragraph("Resumo por Material", h2))
    rm = [["Cód. Material", "Descrição", "Qtd Sistema", "Qtd Física", "% Acuracidade"]]
    for m in resumo_material[:40]:
        rm.append([m["cod_material"], m["descricao"][:40], str(m["qtd_sistema"]), str(m["qtd_fisica"]), f'{m["acuracidade"]}%'])
    rmt = Table(rm, colWidths=[28 * mm, 70 * mm, 24 * mm, 24 * mm, 28 * mm])
    rmt.setStyle(_table_style())
    elems.append(rmt)
    elems.append(Spacer(1, 12))

    # Pendências resumo
    elems.append(Paragraph("Resumo de Pendências", h2))
    pd = [["SKU", "Material", "Status Inv.", "Gatilho", "Resolução", "Status"]]
    for p in pendencias[:40]:
        pd.append([p.get("sku_fisico") or p.get("sku_sistema", ""), p.get("cod_material", ""),
                   p.get("status_inventario", ""), p.get("gatilho", ""), (p.get("resolucao") or "")[:20],
                   p.get("resolucao_status", "")])
    pdt = Table(pd, colWidths=[26 * mm, 26 * mm, 32 * mm, 30 * mm, 30 * mm, 24 * mm])
    pdt.setStyle(_table_style())
    elems.append(pdt)
    elems.append(Spacer(1, 10))
    gerado = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    elems.append(Paragraph(f"<font size=7 color='#6B7280'>Relatório gerado em {gerado} — GAMA Gestão Industrial</font>", normal))

    doc.build(elems)
    return buf.getvalue()


def _table_style():
    return TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
    ])
